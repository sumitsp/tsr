"""Tag the junk subtype of pages a review sheet already called junk.

The review sheet (Excel or CSV) has a folder column, an image id column, a
"Junk Page" column, and a "Blank Page" column. Only the images with Junk Page =
Yes are opened, and each gets its junk type. Images with Blank Page = Yes are
added as BLANK without opening them. Everything goes into the same labels CSV
that main.py writes, and a page that already has a row there is left alone.

Edit RUN CONFIG, then run:  python b_jnk_creator/junk_tagger.py
Open the printed URL. Enter saves the page and opens the next junk image.
"""

from __future__ import annotations

import csv
import logging
import os
import re
import sys
import threading
from dataclasses import dataclass
from pathlib import Path

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from b_jnk_creator.catalog import (
    IMAGE_EXTENSIONS,
    ImageItem,
    _natural_key,
    connect_container,
    read_blob_image,
    read_local_image,
)
from b_jnk_creator.labels import LabelStore
from b_jnk_creator.server import App, serve

# ============================== RUN CONFIG ==============================
# The review sheet: .xlsx or .csv
REVIEW_SHEET = Path(r"C:\Users\sumit.pandey\Desktop\Imaging\junk_review.xlsx")
# Empty means the first sheet of the workbook.
SHEET_NAME = ""

# Column names in the review sheet. Empty means find it by name:
#   folder: "folder name" / "folder"   id: "id" / "image id" / "image name" / "page"
#   junk: the column with "junk" in its name   blank: the column with "blank" in its name
FOLDER_COLUMN = ""
ID_COLUMN = ""
JUNK_COLUMN = ""
BLANK_COLUMN = ""

INPUT_SOURCE = "local"  # "local" or "blob"
# local: one folder or a list of folders. Each is searched at any depth for the chart
# folders the sheet names; a chart's images may sit in it or in its subfolders.
#   LOCAL_INPUT = [Path(r"C:\...\batch1"), Path(r"D:\scans")]
LOCAL_INPUT = Path(r"C:\Users\sumit.pandey\Desktop\Imaging\images")
# blob: PREFIX/<chart folder>/<image>
STORAGE_ACCOUNT = "azsadve2aipoc"
CONTAINER_NAME = "YOUR_CONTAINER_NAME"
PREFIX = "Run1/Batch1/DEID_PNGs/"

# The existing labels CSV. New rows are appended. Close it in Excel while tagging.
CSV_PATH = Path(r"C:\Users\sumit.pandey\Desktop\Imaging\blank_junk_labels.csv")

HOST = "127.0.0.1"
# Not main.py's port, so the two pages remember their places apart.
PORT = 8768
# ========================================================================

LOGGER = logging.getLogger("b_jnk_creator")

YES = {"yes", "y", "true", "1", "x"}
FOLDER_NAMES = ("foldername", "folder", "chartfolder", "chartname", "chart")
ID_NAMES = ("id", "imageid", "imagename", "image", "pageid", "page", "pagenumber", "pageno", "filename")


@dataclass(frozen=True)
class SheetRow:
    folder: str
    image_id: str
    junk: bool
    blank: bool


def _norm(name: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name or "").lower())


def _cell(value: object) -> str:
    """Excel gives 3.0 for an id typed as 3."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value if value is not None else "").strip()


def _read_table(path: Path, sheet: str) -> tuple[list[str], list[list[str]]]:
    if not path.is_file():
        raise FileNotFoundError(f"REVIEW_SHEET does not exist: {path}")
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            table = [row for row in csv.reader(fh)]
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        ws = book[sheet] if sheet else book.worksheets[0]
        table = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        book.close()
    else:
        raise ValueError(f"REVIEW_SHEET must be .xlsx or .csv, got {path.suffix}")
    while table and not any(table[0]):
        table.pop(0)
    if not table:
        raise ValueError(f"{path} is empty")
    return [str(h).strip() for h in table[0]], table[1:]


def _column(headers: list[str], configured: str, exact: tuple[str, ...], contains: str, what: str) -> int:
    normed = [_norm(h) for h in headers]
    if configured:
        if _norm(configured) in normed:
            return normed.index(_norm(configured))
        raise ValueError(f"{what} column {configured!r} is not in the sheet. Columns: {headers}")
    for name in exact:
        if name in normed:
            return normed.index(name)
    if contains:
        hits = [i for i, h in enumerate(normed) if contains in h]
        if len(hits) == 1:
            return hits[0]
    raise ValueError(f"Could not find the {what} column. Set {what.upper()}_COLUMN. Columns: {headers}")


def read_review_sheet(path: Path, sheet: str = "") -> list[SheetRow]:
    headers, table = _read_table(path, sheet)
    folder_i = _column(headers, FOLDER_COLUMN, FOLDER_NAMES, "folder", "folder")
    id_i = _column(headers, ID_COLUMN, ID_NAMES, "", "id")
    junk_i = _column(headers, JUNK_COLUMN, (), "junk", "junk")
    blank_i = _column(headers, BLANK_COLUMN, (), "blank", "blank")
    LOGGER.info(
        "Review sheet columns: folder=%r id=%r junk=%r blank=%r",
        headers[folder_i], headers[id_i], headers[junk_i], headers[blank_i],
    )
    rows: list[SheetRow] = []
    for raw in table:
        get = lambda i: _cell(raw[i]) if i < len(raw) else ""  # noqa: E731
        folder, image_id = get(folder_i), get(id_i)
        if not folder or not image_id:
            continue
        # An id written as a path keeps only the file part.
        image_id = re.split(r"[\\/]", image_id)[-1]
        rows.append(SheetRow(
            folder,
            image_id,
            get(junk_i).lower() in YES,
            get(blank_i).lower() in YES,
        ))
    return rows


def _find_chart_dirs(roots: list[Path], folders: list[str]) -> dict[str, list[Path]]:
    """Chart folder name -> every directory with that name under any root (any depth)."""
    wanted = {f.casefold(): f for f in folders}
    found: dict[str, list[Path]] = {}
    for root in roots:
        if not root.is_dir():
            LOGGER.warning("LOCAL_INPUT folder does not exist: %s", root)
            continue
        LOGGER.info("Searching %s", root)
        for dirpath, dirnames, _ in os.walk(root):
            for name in dirnames:
                key = name.casefold()
                if key in wanted:
                    found.setdefault(wanted[key], []).append(Path(dirpath) / name)
    return found


def _list_folders_local(roots: Path | list[Path], folders: list[str]) -> list[ImageItem]:
    """Images of each chart folder, wherever it is under the roots. The first copy of an
    image name wins when a chart appears in more than one place."""
    roots = [Path(r) for r in (roots if isinstance(roots, (list, tuple)) else [roots])]
    dirs = _find_chart_dirs(roots, folders)
    items: list[ImageItem] = []
    for folder in folders:
        seen: set[str] = set()
        found = []
        for chart_dir in dirs.get(folder, []):
            for path in sorted(chart_dir.rglob("*")):
                if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS and path.name.casefold() not in seen:
                    seen.add(path.name.casefold())
                    found.append(ImageItem(folder, path.name, str(path)))
        if not found:
            LOGGER.warning("No images found for folder %s under LOCAL_INPUT", folder)
        items.extend(sorted(found, key=lambda item: _natural_key(item.image_name)))
    return items


def _list_folders_blob(container, prefix: str, folders: list[str]) -> list[ImageItem]:
    if prefix and not prefix.endswith("/"):
        prefix += "/"
    items: list[ImageItem] = []
    for folder in folders:
        found = []
        for blob in container.list_blobs(name_starts_with=f"{prefix}{folder}/"):
            name = blob.name[len(prefix) + len(folder) + 1:]
            if Path(name).suffix.lower() in IMAGE_EXTENSIONS:
                found.append(ImageItem(folder, name, blob.name))
        if not found:
            LOGGER.warning("No images under %s%s/", prefix, folder)
        items.extend(sorted(found, key=lambda item: _natural_key(Path(item.image_name).name)))
    return items


def match_images(rows: list[SheetRow], images: list[ImageItem]) -> dict[tuple[str, str], ImageItem]:
    """(folder, id) from the sheet -> image. The id may be the file name or its stem."""
    by_key: dict[tuple[str, str], ImageItem] = {}
    for item in images:
        folder = item.folder.casefold()
        name = Path(item.image_name).name.casefold()
        by_key.setdefault((folder, name), item)
        by_key.setdefault((folder, Path(name).stem), item)
    out: dict[tuple[str, str], ImageItem] = {}
    for row in rows:
        item = by_key.get((row.folder.casefold(), row.image_id.casefold()))
        if item is not None:
            out[(row.folder, row.image_id)] = item
    return out


def prepare(rows: list[SheetRow], images: list[ImageItem], store: LabelStore) -> list[ImageItem]:
    """Append the blank pages to the CSV and return the junk images to tag, in folder order."""
    matched = match_images(rows, images)

    blank_rows = []
    blank_missing = 0
    for row in rows:
        if not row.blank or row.junk:
            continue
        item = matched.get((row.folder, row.image_id))
        if item is None:
            blank_missing += 1
            # Still written: the training join only needs the page number in the name.
            blank_rows.append((row.folder, row.image_id, "BLANK", ""))
        else:
            blank_rows.append((item.folder, item.image_name, "BLANK", ""))
    added = store.add_missing(blank_rows)
    LOGGER.info(
        "Blank Page = Yes: %s rows, %s appended, %s already in the CSV (%s with no image found)",
        len(blank_rows), added, len(blank_rows) - added, blank_missing,
    )

    junk: list[ImageItem] = []
    seen: set[tuple[str, str]] = set()
    missing: list[str] = []
    for row in rows:
        if not row.junk:
            continue
        item = matched.get((row.folder, row.image_id))
        if item is None:
            missing.append(f"{row.folder}/{row.image_id}")
            continue
        if (item.folder, item.image_name) not in seen:
            seen.add((item.folder, item.image_name))
            junk.append(item)
    if missing:
        LOGGER.warning("%s Junk Page = Yes rows have no image, e.g. %s", len(missing), ", ".join(missing[:5]))
    folder_order = {}
    for item in junk:
        folder_order.setdefault(item.folder, len(folder_order))
    junk.sort(key=lambda item: (folder_order[item.folder], _natural_key(Path(item.image_name).name)))
    tagged = sum(1 for item in junk if store.has_row(item.folder, item.image_name))
    LOGGER.info("Junk Page = Yes: %s images to tag, %s already tagged", len(junk), tagged)
    return junk


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    source = INPUT_SOURCE.strip().lower()
    if source not in {"local", "blob"}:
        raise SystemExit('INPUT_SOURCE must be "local" or "blob"')
    if source == "blob" and (not CONTAINER_NAME or CONTAINER_NAME == "YOUR_CONTAINER_NAME"):
        raise SystemExit("Set CONTAINER_NAME in b_jnk_creator/junk_tagger.py")

    store = LabelStore(CSV_PATH)
    app = App(store)

    def load() -> None:
        try:
            rows = read_review_sheet(REVIEW_SHEET, SHEET_NAME)
            folders = list(dict.fromkeys(row.folder for row in rows if row.junk or row.blank))
            LOGGER.info("%s sheet rows, %s folders with a junk or blank page", len(rows), len(folders))
            if source == "local":
                app.read_image = read_local_image
                images = _list_folders_local(LOCAL_INPUT, folders)
            else:
                container = connect_container(STORAGE_ACCOUNT, CONTAINER_NAME)
                app.read_image = lambda item: read_blob_image(container, item)
                images = _list_folders_blob(container, PREFIX, folders)
            junk = prepare(rows, images, store)
            if not junk:
                raise FileNotFoundError("No image found for any Junk Page = Yes row")
            app.set_items(junk)
        except Exception as exc:
            LOGGER.exception("Could not load the review sheet")
            app.fail(f"{type(exc).__name__}: {exc}")

    threading.Thread(target=load, name="review-sheet", daemon=True).start()
    LOGGER.info("CSV: %s", CSV_PATH)
    serve(app, HOST, PORT)


if __name__ == "__main__":
    main()
